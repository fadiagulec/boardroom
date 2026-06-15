#!/usr/bin/env python3
import re, glob, os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

INK=RGBColor(0x1a,0x17,0x14); ACCENT=RGBColor(0xc2,0x41,0x0c); ACCENT2=RGBColor(0x0f,0x76,0x6e)
GOLD=RGBColor(0xb0,0x84,0x42); MUTED=RGBColor(0x6b,0x62,0x59)

order=[
 ("Module 0 — Welcome + Your AI Audit","masterclass/scripts/module-00-welcome.md"),
 ("Week 1 — Foundation","masterclass/scripts/module-01-foundation.md"),
 ("Week 2 — Projects","masterclass/scripts/module-02-projects.md"),
 ("Week 3 — Content","masterclass/scripts/module-03-content-engine.md"),
 ("Week 4 — Cowork","masterclass/scripts/module-04-cowork.md"),
 ("Week 5 — Claude Code (Finale)","masterclass/scripts/module-05-code.md"),
 ("Bonus — Claude Essentials","masterclass/scripts/bonus-foundations.md"),
]

def clean(t):
    t=re.sub(r"\*\*(.+?)\*\*",r"\1",t)
    t=re.sub(r"\*(.+?)\*",r"\1",t)
    t=t.replace("`","")
    return t.strip()

doc=Document()
# base style
st=doc.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(11)

# COVER
doc.add_paragraph()
h=doc.add_paragraph(); h.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=h.add_run("BUILT IN A DAY"); r.bold=True; r.font.size=Pt(40); r.font.color.rgb=INK
s=doc.add_paragraph(); s.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=s.add_run("Full Training Material — what you teach in every lesson"); r.font.size=Pt(15); r.font.color.rgb=MUTED
s=doc.add_paragraph(); s.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=s.add_run("Master Claude · Build apps, agents & automations · 5 weeks"); r.font.size=Pt(12); r.font.color.rgb=ACCENT
note=doc.add_paragraph(); note.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=note.add_run("Read the NARRATION aloud = the words you teach.  SLIDES = what's on screen.  ON SCREEN (DEMO) = switch to Claude and do it live."); r.italic=True; r.font.size=Pt(10); r.font.color.rgb=MUTED
doc.add_page_break()

for secttitle, path in order:
    if not os.path.exists(path): continue
    # Section header
    p=doc.add_paragraph(); r=p.add_run(secttitle); r.bold=True; r.font.size=Pt(26); r.font.color.rgb=ACCENT
    lines=open(path,encoding="utf-8").read().splitlines()
    incode=False
    for ln in lines:
        s=ln.rstrip()
        if s.startswith("# "):  # file title, skip (covered by section header)
            continue
        if s.startswith("```"):
            incode=not incode
            if incode:
                lab=doc.add_paragraph(); r=lab.add_run("SLIDES (put on screen):"); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=ACCENT2
            continue
        if incode:
            cp=doc.add_paragraph(); cp.paragraph_format.left_indent=Inches(0.3)
            r=cp.add_run(s); r.font.name="Consolas"; r.font.size=Pt(10); r.font.color.rgb=INK
            continue
        if not s.strip():
            continue
        if s.startswith("## "):
            p=doc.add_paragraph(); p.space_before=Pt(10)
            r=p.add_run(clean(s[3:])); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=INK
        elif s.startswith("### "):
            p=doc.add_paragraph(); r=p.add_run(clean(s[4:])); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=INK
        elif s.startswith("> "):
            body=clean(s[2:])
            if not body: continue
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.25)
            r=p.add_run(body); r.font.size=Pt(12)
            if body.lower().startswith(("note","resource","provide","bonus","phase")):
                r.italic=True; r.font.color.rgb=MUTED; r.font.size=Pt(10.5)
        elif s.startswith("NARRATION"):
            p=doc.add_paragraph(); r=p.add_run("SAY THIS:"); r.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=ACCENT2
        elif s.startswith("[DEMO]"):
            p=doc.add_paragraph(); r=p.add_run("ON SCREEN — "+clean(s[6:]).lstrip(": ")); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=ACCENT
        elif s.startswith("✅"):
            p=doc.add_paragraph(style="List Bullet"); r=p.add_run(clean(s.replace("✅","Action:",1))); r.font.size=Pt(11)
        elif s.startswith("- ") or s.startswith("* "):
            p=doc.add_paragraph(style="List Bullet"); r=p.add_run(clean(s[2:])); r.font.size=Pt(11)
        elif s.startswith("Resource"):
            p=doc.add_paragraph(); r=p.add_run(clean(s)); r.italic=True; r.font.size=Pt(10); r.font.color.rgb=MUTED
        elif s=="---":
            continue
        else:
            p=doc.add_paragraph(); r=p.add_run(clean(s)); r.font.size=Pt(11)
    doc.add_page_break()

doc.save("masterclass/Built-in-a-Day-TRAINING-MATERIAL.docx")
print("saved manual")
