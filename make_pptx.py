#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

INK=RGBColor(0x1a,0x17,0x14); BG=RGBColor(0xf6,0xf3,0xee); CREAM=RGBColor(0xf6,0xf3,0xee)
ACCENT=RGBColor(0xc2,0x41,0x0c); ACCENT2=RGBColor(0x0f,0x76,0x6e); GOLD=RGBColor(0xb0,0x84,0x42)
MUTED=RGBColor(0x6b,0x62,0x59); CARDLINE=RGBColor(0x3a,0x2e,0x22); WHITE=RGBColor(0xff,0xff,0xff)
DARKCREAM=RGBColor(0xd8,0xcf,0xc2)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
W=13.333; H=7.5

def slide(bg):
    s=prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb=bg
    return s

def box(s,left,top,width,height,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height))
    tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    return tf

def line(tf,text,size,color,bold=False,align=PP_ALIGN.LEFT,space_after=6,first=False):
    p=tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment=align; p.space_after=Pt(space_after)
    r=p.add_run(); r.text=text; f=r.font
    f.size=Pt(size); f.bold=bold; f.color.rgb=color; f.name="Calibri"
    return p

def kicker(s,text,color=ACCENT,top=0.55):
    tf=box(s,0.9,top,11.5,0.5)
    line(tf,text.upper(),14,color,bold=True,first=True)

def title(s,text,color=INK,top=1.0,size=40):
    tf=box(s,0.9,top,11.6,1.6)
    line(tf,text,size,color,bold=True,first=True)

def bullets(s,items,top,color=INK,size=18,left=0.9,width=11.6,height=4.6,gap=8):
    tf=box(s,left,top,width,height)
    for i,(txt,b) in enumerate(items):
        line(tf,txt,size,color,bold=b,first=(i==0),space_after=gap)

# 1 TITLE
s=slide(INK)
kicker(s,"The Masterclass",GOLD,0.9)
title(s,"BUILT IN A DAY",CREAM,1.35,60)
tf=box(s,0.9,2.7,11.6,0.7); line(tf,"Master Claude · Build apps, agents & automations · 5 weeks",24,DARKCREAM,first=True)
tf=box(s,0.9,3.6,11.6,2.2)
line(tf,'"I\'m 55. I spent 30 years in corporate with no tech background. Six months ago AI scared me — so I learned it. I got hooked on Claude and built a system that runs a 300-vehicle fleet in ONE day. If someone with zero tech background can do this — anyone can."',18,CREAM,first=True)
tf=box(s,0.9,6.4,11.6,0.6); line(tf,"🎥 Video-first    🟢 Live weekly (Zoom)    🖥️ Real Claude on screen    🏫 Teachable",15,GOLD,first=True)

# 2 PROMISE
s=slide(BG); kicker(s,"The promise"); title(s,'From "I\'ve heard of Claude" to building real things')
bullets(s,[
 ("🎯  Who it's for — professionals who feel behind on AI and want to build with it, not be replaced by it. No coding background needed.",False),
 ("⚡  The format — short full-walkthrough videos, a weekly live build call, fill-in workbooks, a copy-paste prompt library.",False),
 ("🏆  Proof — the one-day 300-vehicle fleet app · ~10 apps in 6 months · a real before/after AI Audit jump.",False),
 ("🧭  A do-it course, not an information dump — every week you build something real.",False),
],2.5,size=20,gap=14)

# 3 ALL LEVELS
s=slide(BG); kicker(s,"Built for all levels — friendly by design"); title(s,"Nobody gets left behind")
bullets(s,[
 ("Total beginners follow every click — nothing assumed, plain English, a 'do this now' after each step.",False),
 ("Already using Claude? A 60-second 'fast-pass' card at the top of each module so you're never bored.",False),
 ("Tone: warm, encouraging, zero jargon. Every hard moment has a 'stuck? paste the error and ask Claude' safety net.",False),
 ("The vibe: a friend who figured it out, showing you how — exactly the founder's story.",False),
],2.6,size=20,gap=16)

# 4 JOURNEY
s=slide(BG); kicker(s,"The journey"); title(s,"One welcome + five weekly phases + a bonus library")
bullets(s,[
 ("Week 1 — Foundation   ·   set Claude up properly   ·   2–3 hrs",True),
 ("Week 2 — Projects   ·   an army of specialised advisors   ·   3–4 hrs",True),
 ("Week 3 — Content   ·   research → strategy → 30-day plan   ·   4–5 hrs",True),
 ("Week 4 — Cowork   ·   hand over whole workflows   ·   3–4 hrs",True),
 ("Week 5 — Claude Code (FINALE)   ·   build & publish a real tool   ·   4–5 hrs",True),
 ("+ Bonus — Claude Essentials   ·   10 short reference videos",False),
],2.5,size=20,gap=14)

# Module slides
def module(num,kick,ttl,dur,vids):
    s=slide(BG)
    kicker(s,kick)
    tf=box(s,0.9,1.0,9.2,1.4); line(tf,f"{num}.  {ttl}",34,INK,bold=True,first=True)
    tf=box(s,10.3,1.05,2.2,0.6); line(tf,dur,16,ACCENT2,bold=True,first=True,align=PP_ALIGN.RIGHT)
    bullets(s,[(v,False) for v in vids],2.5,size=18,gap=10)

module("0","Welcome","Welcome + Your AI Audit","~15 min",[
 "V0.1  Why this exists — your hero story (on camera)",
 "V0.2  How the 5 weeks work — the roadmap",
 "V0.3  Take your AI Audit — record your /40 baseline"])
module("1","Week 1 · Setup","Foundation","2–3 hrs",[
 "V1.1  Get set up — Pro + desktop app + first message",
 "V1.2  The setting everyone skips — Personal Preferences",
 "V1.3  Privacy + your name",
 "V1.4  The 4 modes — Chat · Projects · Cowork · Code"])
module("2","Week 2 · Foundations","Projects — an army of advisors","3–4 hrs",[
 "V2.1  Why one generic chatbot is the problem",
 "V2.2  Design your project architecture",
 "V2.3  The master workspace prompt (live build)",
 "V2.4  The handover — never lose context",
 "V2.5  Checkpoint + test it"])
module("3","Week 3 · Content (biggest)","Strategist + creative partner","4–5 hrs",[
 "V3.1  Why strategy first — the workflow",
 "V3.2  Write your audience research brief",
 "V3.3  Run across tools + synthesise",
 "V3.4  Brain dump + belief interview",
 "V3.5  Generate your strategy (Opus)",
 "V3.6  Plan 30 days of content",
 "V3.7  Week 3 checkpoint"])
module("4","Week 4 · Cowork","Hand over the work","3–4 hrs",[
 "V4.1  Chat vs Cowork = agency",
 "V4.2  Set up the workspace + guardrails",
 "V4.3  Give Cowork a brain — CLAUDE.md + MEMORY.md",
 "V4.4  Plan the daily brief (in Chat)",
 "V4.5  Run, fix, schedule it",
 "V4.6  Week 4 check-in"])
module("5","Week 5 · Build · FINALE","Claude Code — build a real tool","4–5 hrs",[
 "V5.1  What Claude Code is",
 "V5.2  Install + Antigravity",
 "V5.3  CLAUDE.md + permissions",
 "V5.4  Spec it first — the PRD",
 "V5.5  Build it — Plan mode → build",
 "V5.6  Save & publish (GitHub + Vercel)",
 "V5.7  You did it — re-take your AI Audit"])

# BONUS
s=slide(BG); kicker(s,"Bonus · Reference"); title(s,"Claude Essentials — 10 short videos")
bullets(s,[
 ("B1 Chat vs Cowork vs Code    ·    B2 Which model (Haiku/Sonnet/Opus)",False),
 ("B3 What Claude can do (Artifacts)    ·    B4 Why Claude over other LLMs",False),
 ("B5 Prompt engineering core    ·    B6 Prompt chaining",False),
 ("B7 Skills deep dive    ·    B8 Projects deep dive",False),
 ("B9 Connectors / MCP    ·    B10 The full stack together",False),
],2.7,size=20,gap=16)

# HOSTING
s=slide(INK); kicker(s,"How it's delivered",GOLD); title(s,"Hosting & the live layer",CREAM)
bullets(s,[
 ("🏫  Home: Teachable — sellable course, native video hosting, payments, drip weekly. No community to manage.",False),
 ("🟢  Live: Zoom — one weekly call (recap → live build → hot seats → Q&A). Recorded → uploaded as a bonus lesson.",False),
 ("🎥  Record: Loom — talking head + screen in one click, your face over the live Claude demo. (Screen Studio for polish.)",False),
],2.7,color=CREAM,size=20,gap=18)

# RULES
s=slide(BG); kicker(s,"How the videos feel"); title(s,"Three production rules")
bullets(s,[
 ("Real Claude only — every demo is the actual Claude product doing the actual thing. Never stock footage or mockups.",False),
 ("Show everything — full uncut walkthroughs, every click, real output. Slides are bookends; the screen is the lesson.",False),
 ("Interactive — pause-and-do cards, fill-in workbooks, end-of-module checkpoints, weekly live calls.",False),
],2.7,size=20,gap=18)

# PRODUCTION
s=slide(BG); kicker(s,"How it gets made"); title(s,"The production line")
bullets(s,[
 ("1.  Script ✅ (already written)",True),
 ("2.  Record talking-head + screen in Loom (real Claude)",True),
 ("3.  Tidy in Descript — optional",True),
 ("4.  Upload to Teachable → drip one module/week",True),
 ("Flow:  Script → Loom → (Descript) → Teachable → Zoom live calls",False),
],2.7,size=20,gap=16)

# CLOSE
s=slide(INK); kicker(s,"Built in a Day",GOLD,1.4)
tf=box(s,0.9,2.2,11.6,2.6)
line(tf,"If someone with zero tech background",44,CREAM,bold=True,first=True,space_after=2)
line(tf,"can do this — anyone can.",44,CREAM,bold=True,space_after=10)
tf=box(s,0.9,5.0,11.6,0.7); line(tf,"Let me show you exactly how.",24,GOLD,first=True)
tf=box(s,0.9,6.6,11.6,0.5); line(tf,"Built in a Day · by Fadia",13,DARKCREAM,first=True)

prs.save("masterclass/Built-in-a-Day.pptx")
print("saved", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
